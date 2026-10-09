import csv
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "labels" / "manifest.csv"

ws = load_workbook(ROOT / "Paintings_Metadata.xlsx", read_only=True)["Paintings"]
rows = list(ws.iter_rows(values_only=True))
h = {k: i for i, k in enumerate(rows[0])}
OUT.parent.mkdir(exist_ok=True)
with OUT.open("w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["id", "file", "known_artist", "known_title", "known_year"])
    for r in rows[1:]:
        path = Path(r[h["Folder"]]) / r[h["Image file"]]
        assert (ROOT / path).exists(), path
        w.writerow([r[h["ID"]], path.as_posix(), r[h["Artist (EN)"]] or "", r[h["Title (EN)"]] or "", r[h["Year"]] or ""])
print(OUT, len(rows) - 1)
