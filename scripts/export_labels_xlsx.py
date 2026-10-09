import csv, json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
L = ROOT / "labels"
OUT = L / "labels_review.xlsx"

recs = [json.loads(l) for l in (L / "labels.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
with (L / "splits.csv").open(encoding="utf-8") as f:
    split = {r["id"]: r["split"] for r in csv.DictReader(f)}

cols = [("ID", 9), ("Split", 7), ("Image", 12), ("Known artist", 24), ("Known title", 24), ("Known year", 9),
        ("Description", 70), ("Theme", 28), ("Mood", 20), ("Objects", 40), ("Style guess", 28),
        ("Medium guess", 26), ("Era guess", 22), ("Lighting", 26), ("Composition", 36), ("Title guess", 26)]
cols += [(f"Color {i}", 14) for i in range(1, 7)]

wb = Workbook()
ws = wb.active
ws.title = "Labels"
ws.append([c for c, _ in cols])
for r in recs:
    pal = r["palette"] or []
    row = [r["id"], split.get(r["id"], ""), "open", r["known_artist"], r["known_title"], r["known_year"],
           r["description"], r["theme"], r["mood"], ", ".join(r["objects"] or []), r["style_guess"],
           r["medium_guess"], r["era_guess"], r["lighting"], r["composition"], r["title_guess"]]
    row += [f"{c['hex']} {c['pct']}%" for c in pal] + [""] * (6 - len(pal))
    ws.append(row)
    n = ws.max_row
    img = ws.cell(n, 3)
    img.hyperlink = str(ROOT / r["file"])
    img.font = Font(color="0563C1", underline="single")
    for i, c in enumerate(pal):
        cell = ws.cell(n, 17 + i)
        h = c["hex"].lstrip("#")
        cell.fill = PatternFill("solid", fgColor=h)
        lum = 0.299 * int(h[0:2], 16) + 0.587 * int(h[2:4], 16) + 0.114 * int(h[4:6], 16)
        cell.font = Font(color="FFFFFF" if lum < 128 else "000000")

for i, (_, w) in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDE5F0")
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "B2"
ws.auto_filter.ref = ws.dimensions
wb.save(OUT)
print(OUT, len(recs))
